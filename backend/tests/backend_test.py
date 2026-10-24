"""Backend API tests for Fleet Manager (APR MENSUR DULJEVIĆ)"""
import os
import pytest
import requests
from datetime import date, timedelta

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://maintenance-log-120.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"
ADMIN_EMAIL = "duljevic1366@gmail.com"
ADMIN_PASSWORD = "fleet2026"


@pytest.fixture(scope="session")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, r.text
    data = r.json()
    assert "token" in data and data["email"] == ADMIN_EMAIL
    return data["token"]


@pytest.fixture(scope="session")
def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


# ---------- Auth ----------
class TestAuth:
    def test_login_invalid(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": "wrong"})
        assert r.status_code == 401

    def test_me_requires_auth(self):
        r = requests.get(f"{API}/auth/me")
        assert r.status_code == 401

    def test_me_with_token(self, auth_headers):
        r = requests.get(f"{API}/auth/me", headers=auth_headers)
        assert r.status_code == 200
        assert r.json()["email"] == ADMIN_EMAIL

    def test_vehicles_unauth(self):
        assert requests.get(f"{API}/vehicles").status_code == 401

    def test_dashboard_unauth(self):
        assert requests.get(f"{API}/dashboard").status_code == 401

    def test_word_unauth(self):
        assert requests.get(f"{API}/reports/word").status_code == 401


# ---------- Vehicles CRUD ----------
class TestVehicles:
    def test_list(self, auth_headers):
        r = requests.get(f"{API}/vehicles", headers=auth_headers)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_crud_and_doc_status(self, auth_headers):
        expired = (date.today() - timedelta(days=10)).isoformat()
        soon = (date.today() + timedelta(days=15)).isoformat()
        valid = (date.today() + timedelta(days=200)).isoformat()
        payload = {
            "plate": "TEST-001", "make": "MAN", "model": "TGX", "year": 2020,
            "vehicle_type": "Kamion", "vin": "TESTVIN123",
            "tech_6m": {"expiry": expired, "note": "t"},
            "registration_12m": {"expiry": soon, "note": ""},
            "cmr_insurance": {"expiry": valid, "note": ""},
            "kasko_insurance": {"expiry": None, "note": ""},
        }
        r = requests.post(f"{API}/vehicles", json=payload, headers=auth_headers)
        assert r.status_code == 200, r.text
        v = r.json()
        vid = v["id"]
        assert v["plate"] == "TEST-001"
        assert v["tech_6m"]["status"] == "expired"
        assert v["registration_12m"]["status"] == "expiring"
        assert v["cmr_insurance"]["status"] == "valid"
        assert v["kasko_insurance"]["status"] == "none"

        # GET
        r = requests.get(f"{API}/vehicles/{vid}", headers=auth_headers)
        assert r.status_code == 200 and r.json()["plate"] == "TEST-001"

        # PUT
        payload["make"] = "Volvo"
        r = requests.put(f"{API}/vehicles/{vid}", json=payload, headers=auth_headers)
        assert r.status_code == 200 and r.json()["make"] == "Volvo"

        # Maintenance add
        mp = {"vehicle_id": vid, "date": date.today().isoformat(),
              "service_type": "Zamena ulja", "description": "Test", "cost": 15000}
        rm = requests.post(f"{API}/maintenance", json=mp, headers=auth_headers)
        assert rm.status_code == 200
        mid = rm.json()["id"]

        rlist = requests.get(f"{API}/maintenance", params={"vehicle_id": vid}, headers=auth_headers)
        assert rlist.status_code == 200 and any(m["id"] == mid for m in rlist.json())

        # Delete vehicle (cascade maintenance)
        rd = requests.delete(f"{API}/vehicles/{vid}", headers=auth_headers)
        assert rd.status_code == 200
        assert requests.get(f"{API}/vehicles/{vid}", headers=auth_headers).status_code == 404


# ---------- Dashboard ----------
class TestDashboard:
    def test_dashboard(self, auth_headers):
        r = requests.get(f"{API}/dashboard", headers=auth_headers)
        assert r.status_code == 200
        d = r.json()
        for k in ["total_vehicles", "expired_docs", "expiring_docs", "total_services", "total_cost", "alerts"]:
            assert k in d
        # alerts sorted: expired first
        statuses = [a["status"] for a in d["alerts"]]
        if "expired" in statuses and "expiring" in statuses:
            assert statuses.index("expired") < statuses.index("expiring")


# ---------- Word export ----------
class TestWord:
    def test_word_download(self, auth_headers):
        r = requests.get(f"{API}/reports/word", headers=auth_headers)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert r.content[:2] == b"PK"  # docx is zip
        assert len(r.content) > 1000


# ---------- New in iteration 2: per-vehicle word, documents upload, reminders, cron ----------
WEBHOOK_CRON_SECRET = "c8f4a2e97b1d6350af92e4c71b8d05fa3e6217c9b4d80f5a2e1c7936bd4a0f88"


class TestPerVehicleWord:
    def test_word_per_vehicle(self, auth_headers):
        # Need a vehicle - create ephemeral
        payload = {"plate": "TEST-WORD-1", "make": "MAN", "model": "TGX", "year": 2021,
                   "vehicle_type": "Tegljač", "vin": "WVIN1",
                   "tech_6m": {"expiry": None, "note": ""},
                   "registration_12m": {"expiry": None, "note": ""},
                   "cmr_insurance": {"expiry": None, "note": ""},
                   "kasko_insurance": {"expiry": None, "note": ""}}
        r = requests.post(f"{API}/vehicles", json=payload, headers=auth_headers)
        assert r.status_code == 200
        vid = r.json()["id"]
        try:
            rw = requests.get(f"{API}/reports/word/{vid}", headers=auth_headers)
            assert rw.status_code == 200
            assert rw.content[:2] == b"PK"
            assert rw.headers.get("content-type", "").startswith(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            assert len(rw.content) > 1000
        finally:
            requests.delete(f"{API}/vehicles/{vid}", headers=auth_headers)


class TestDocuments:
    def test_upload_list_download_delete(self, auth_headers, token):
        # 1x1 PNG bytes
        png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
               b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf"
               b"\xc0\x00\x00\x00\x03\x00\x01v\xb5\x8b\xaa\x00\x00\x00\x00IEND\xaeB`\x82")
        # create vehicle
        payload = {"plate": "TEST-DOC-1", "make": "V", "model": "M", "year": 2020,
                   "vehicle_type": "Kamion", "vin": "DVIN1",
                   "tech_6m": {"expiry": None, "note": ""},
                   "registration_12m": {"expiry": None, "note": ""},
                   "cmr_insurance": {"expiry": None, "note": ""},
                   "kasko_insurance": {"expiry": None, "note": ""}}
        rv = requests.post(f"{API}/vehicles", json=payload, headers=auth_headers)
        vid = rv.json()["id"]
        try:
            # invalid ext
            bad = requests.post(f"{API}/vehicles/{vid}/documents",
                                files={"file": ("evil.exe", b"MZ", "application/octet-stream")},
                                data={"category": "ostalo"},
                                headers=auth_headers)
            assert bad.status_code == 400

            # good upload
            up = requests.post(f"{API}/vehicles/{vid}/documents",
                               files={"file": ("test.png", png, "image/png")},
                               data={"category": "registracija"},
                               headers=auth_headers)
            assert up.status_code == 200, up.text
            doc = up.json()
            assert doc["category"] == "registracija"
            assert doc["original_filename"] == "test.png"
            did = doc["id"]

            # list
            lst = requests.get(f"{API}/vehicles/{vid}/documents", headers=auth_headers).json()
            assert any(d["id"] == did for d in lst)

            # download with query auth
            dl = requests.get(f"{API}/documents/{did}/download", params={"auth": token})
            assert dl.status_code == 200
            assert dl.content[:4] == b"\x89PNG"

            # download unauth
            dlu = requests.get(f"{API}/documents/{did}/download")
            assert dlu.status_code == 401

            # delete
            d = requests.delete(f"{API}/documents/{did}", headers=auth_headers)
            assert d.status_code == 200
            lst2 = requests.get(f"{API}/vehicles/{vid}/documents", headers=auth_headers).json()
            assert not any(x["id"] == did for x in lst2)
        finally:
            requests.delete(f"{API}/vehicles/{vid}", headers=auth_headers)


class TestReminders:
    def test_preview(self, auth_headers):
        r = requests.get(f"{API}/reminders/preview", headers=auth_headers)
        assert r.status_code == 200
        d = r.json()
        assert "count" in d and "recipients" in d and "items" in d
        assert isinstance(d["recipients"], list)
        assert isinstance(d["count"], int)

    def test_preview_unauth(self):
        assert requests.get(f"{API}/reminders/preview").status_code == 401

    def test_cron_unauth(self):
        r = requests.post(f"{API}/cron/reminders")
        assert r.status_code == 401

    def test_cron_wrong_secret(self):
        r = requests.post(f"{API}/cron/reminders",
                          headers={"Authorization": "Bearer wrong"})
        assert r.status_code == 401

    def test_cron_ok(self):
        r = requests.post(f"{API}/cron/reminders",
                          headers={"Authorization": f"Bearer {WEBHOOK_CRON_SECRET}"})
        assert r.status_code == 200
        assert r.json().get("status") in ("accepted", "duplicate")


# ---------- Iteration 3: settings recipients + excel ----------
class TestSettingsRecipients:
    def test_get_recipients(self, auth_headers):
        r = requests.get(f"{API}/settings/recipients", headers=auth_headers)
        assert r.status_code == 200
        d = r.json()
        assert isinstance(d.get("recipients"), list)

    def test_unauth(self):
        assert requests.get(f"{API}/settings/recipients").status_code == 401
        assert requests.put(f"{API}/settings/recipients", json={"recipients": []}).status_code == 401

    def test_put_invalid_email(self, auth_headers):
        r = requests.put(f"{API}/settings/recipients",
                         json={"recipients": ["not-an-email"]},
                         headers=auth_headers)
        assert r.status_code == 422

    def test_put_and_get_roundtrip_and_reset(self, auth_headers):
        # Save original
        orig = requests.get(f"{API}/settings/recipients", headers=auth_headers).json()["recipients"]
        try:
            new_list = ["duljevic1366@gmail.com", "test_extra@example.com"]
            r = requests.put(f"{API}/settings/recipients", json={"recipients": new_list}, headers=auth_headers)
            assert r.status_code == 200
            assert sorted(r.json()["recipients"]) == sorted([e.lower() for e in new_list])
            # GET reflects
            g = requests.get(f"{API}/settings/recipients", headers=auth_headers).json()
            assert sorted(g["recipients"]) == sorted([e.lower() for e in new_list])
            # /reminders/preview recipients match
            prev = requests.get(f"{API}/reminders/preview", headers=auth_headers).json()
            assert sorted(prev["recipients"]) == sorted([e.lower() for e in new_list])
        finally:
            # RESET
            reset = orig if orig else ["duljevic1366@gmail.com"]
            requests.put(f"{API}/settings/recipients", json={"recipients": reset}, headers=auth_headers)


class TestExcelExport:
    def test_excel_unauth(self):
        assert requests.get(f"{API}/reports/excel").status_code == 401

    def test_excel_download(self, auth_headers):
        r = requests.get(f"{API}/reports/excel", headers=auth_headers)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert r.content[:2] == b"PK"  # xlsx is zip
        assert len(r.content) > 500
        # Verify two sheets
        import io as _io
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(r.content), read_only=True)
        assert "Vozila" in wb.sheetnames
        assert "Održavanje" in wb.sheetnames



# ---------- Iteration 6: change password ----------
class TestChangePassword:
    def test_unauth(self):
        r = requests.post(f"{API}/auth/change-password",
                          json={"current_password": "x", "new_password": "yyyyyy"})
        assert r.status_code == 401

    def test_full_change_flow(self, auth_headers):
        temp_pw = "Temp#2026x"
        # wrong current
        r = requests.post(f"{API}/auth/change-password",
                          json={"current_password": "wrongpw", "new_password": temp_pw},
                          headers=auth_headers)
        assert r.status_code == 400
        assert "Trenutna lozinka nije tačna" in r.json().get("detail", "")

        # too short
        r = requests.post(f"{API}/auth/change-password",
                          json={"current_password": ADMIN_PASSWORD, "new_password": "abc"},
                          headers=auth_headers)
        assert r.status_code == 400
        assert "6 karaktera" in r.json().get("detail", "")

        # valid change
        r = requests.post(f"{API}/auth/change-password",
                          json={"current_password": ADMIN_PASSWORD, "new_password": temp_pw},
                          headers=auth_headers)
        assert r.status_code == 200, r.text

        try:
            # login with new password works
            lr = requests.post(f"{API}/auth/login",
                               json={"email": ADMIN_EMAIL, "password": temp_pw})
            assert lr.status_code == 200
            new_token = lr.json()["token"]

            # login with old password fails
            old = requests.post(f"{API}/auth/login",
                                json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
            assert old.status_code == 401
        finally:
            # CRITICAL: change back to original
            new_headers = {"Authorization": f"Bearer {new_token}"}
            rb = requests.post(f"{API}/auth/change-password",
                               json={"current_password": temp_pw, "new_password": ADMIN_PASSWORD},
                               headers=new_headers)
            assert rb.status_code == 200, rb.text
            # verify original works
            final = requests.post(f"{API}/auth/login",
                                  json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
            assert final.status_code == 200
