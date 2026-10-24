"""Iteration 16 - Tahograf (vehicle) + Kartica vozača (driver) end-to-end backend tests.

Run sequentially (module-scoped state used across tests).
"""
import os
import io
import zipfile
import datetime as dt
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://maintenance-log-120.preview.emergentagent.com").rstrip("/")
ADMIN = {"email": "duljevic1366@gmail.com", "password": "fleet2026"}

STATE = {"vehicle_id": None, "driver_id": None}


def _iso(days):
    return (dt.date.today() + dt.timedelta(days=days)).isoformat()


def _docx_text(content: bytes) -> str:
    """Extract text from all word/*.xml files inside a .docx"""
    out = []
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for name in zf.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                try:
                    out.append(zf.read(name).decode("utf-8", errors="ignore"))
                except Exception:
                    pass
    return "\n".join(out)


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json=ADMIN, timeout=30)
    assert r.status_code == 200, r.text
    tok = r.json().get("token")
    if tok:
        s.headers.update({"Authorization": f"Bearer {tok}"})
    yield s


# ============ VEHICLE tachograph ============
def test_01_vehicle_create_with_tachograph(client):
    payload = {
        "plate": "TESTTAH16",
        "brand": "Volvo",
        "model": "FH",
        "odometer_km": 1000,
        "tachograph": {"expiry": _iso(5)},
    }
    r = client.post(f"{BASE_URL}/api/vehicles", json=payload, timeout=30)
    assert r.status_code in (200, 201), r.text
    v = r.json()
    vid = v.get("id") or v.get("_id")
    assert vid
    STATE["vehicle_id"] = vid
    assert v["tachograph"]["expiry"] == _iso(5)
    assert v["tachograph"].get("status") in ("expiring", "expired", "valid", "none")


def test_02_vehicle_get_returns_tachograph(client):
    vid = STATE["vehicle_id"]
    r = client.get(f"{BASE_URL}/api/vehicles/{vid}", timeout=30)
    assert r.status_code == 200
    v = r.json()
    assert v["tachograph"]["expiry"] == _iso(5)
    assert v["tachograph"]["status"] == "expiring"


def test_03_vehicle_put_updates_tachograph(client):
    vid = STATE["vehicle_id"]
    body = {"plate": "TESTTAH16", "brand": "Volvo", "model": "FH",
            "odometer_km": 1000, "tachograph": {"expiry": _iso(400)}}
    r = client.put(f"{BASE_URL}/api/vehicles/{vid}", json=body, timeout=30)
    assert r.status_code == 200, r.text
    assert r.json()["tachograph"]["expiry"] == _iso(400)
    assert r.json()["tachograph"]["status"] == "valid"


def test_04_dashboard_lists_tachograph_alert(client):
    vid = STATE["vehicle_id"]
    body = {"plate": "TESTTAH16", "brand": "Volvo", "model": "FH",
            "odometer_km": 1000, "tachograph": {"expiry": _iso(10)}}
    client.put(f"{BASE_URL}/api/vehicles/{vid}", json=body, timeout=30)
    r = client.get(f"{BASE_URL}/api/dashboard", timeout=30)
    assert r.status_code == 200
    text = str(r.json())
    assert "Tahograf" in text, f"Tahograf not in dashboard payload: {text[:600]}"


# ============ DRIVER driver_card ============
def test_05_driver_create_with_driver_card(client):
    payload = {
        "first_name": "TESTDR16",
        "last_name": "TAHO",
        "phone": "060",
        "driver_card": {"expiry": _iso(8)},
    }
    r = client.post(f"{BASE_URL}/api/drivers", json=payload, timeout=30)
    assert r.status_code in (200, 201), r.text
    d = r.json()
    did = d.get("id") or d.get("_id")
    assert did
    STATE["driver_id"] = did
    assert d["driver_card"]["expiry"] == _iso(8)
    assert d["driver_card"].get("status") == "expiring"


def test_06_driver_get_returns_driver_card(client):
    did = STATE["driver_id"]
    r = client.get(f"{BASE_URL}/api/drivers/{did}", timeout=30)
    assert r.status_code == 200
    assert r.json()["driver_card"]["expiry"] == _iso(8)


def test_07_driver_put_updates_driver_card(client):
    did = STATE["driver_id"]
    body = {"first_name": "TESTDR16", "last_name": "TAHO", "phone": "060",
            "driver_card": {"expiry": _iso(500)}}
    r = client.put(f"{BASE_URL}/api/drivers/{did}", json=body, timeout=30)
    assert r.status_code == 200, r.text
    assert r.json()["driver_card"]["expiry"] == _iso(500)
    assert r.json()["driver_card"]["status"] == "valid"


def test_08_dashboard_lists_driver_card_alert(client):
    did = STATE["driver_id"]
    body = {"first_name": "TESTDR16", "last_name": "TAHO", "phone": "060",
            "driver_card": {"expiry": _iso(6)}}
    client.put(f"{BASE_URL}/api/drivers/{did}", json=body, timeout=30)
    r = client.get(f"{BASE_URL}/api/dashboard", timeout=30)
    assert r.status_code == 200
    text = str(r.json())
    assert "Kartica vozača" in text, f"Kartica vozača missing in dashboard: {text[:600]}"


# ============ EXPORTS ============
def test_09_excel_vehicles_has_columns(client):
    r = client.get(f"{BASE_URL}/api/reports/excel", timeout=60)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert "Vozila" in wb.sheetnames
    headers_v = [c.value for c in wb["Vozila"][1]]
    assert "Tahograf" in headers_v, f"Tahograf missing: {headers_v}"
    assert "Vozači" in wb.sheetnames
    headers_d = [c.value for c in wb["Vozači"][1]]
    assert "Kartica vozača" in headers_d, f"Kartica vozača missing: {headers_d}"


def test_10_excel_drivers_has_driver_card(client):
    r = client.get(f"{BASE_URL}/api/reports/drivers/excel", timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    sh = wb.active
    headers = [c.value for c in sh[1]]
    assert "Kartica vozača" in headers, headers


def test_11_word_vehicle_contains_tachograph(client):
    vid = STATE["vehicle_id"]
    r = client.get(f"{BASE_URL}/api/reports/word/{vid}", timeout=60)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    text = _docx_text(r.content)
    assert "Tahograf" in text


def test_12_word_fleet_contains_tachograph(client):
    r = client.get(f"{BASE_URL}/api/reports/word", timeout=60)
    assert r.status_code == 200
    text = _docx_text(r.content)
    assert "Tahograf" in text


def test_13_word_drivers_summary_contains_driver_card(client):
    r = client.get(f"{BASE_URL}/api/reports/word/drivers", timeout=60)
    assert r.status_code == 200
    text = _docx_text(r.content)
    assert "Kartica vozača" in text


def test_14_word_driver_single_contains_driver_card(client):
    did = STATE["driver_id"]
    r = client.get(f"{BASE_URL}/api/reports/word/driver/{did}", timeout=60)
    assert r.status_code == 200
    text = _docx_text(r.content)
    assert "Kartica vozača" in text


# ============ REGRESSIONS ============
def test_15_reminders_preview_regression(client):
    r = client.get(f"{BASE_URL}/api/reminders/preview", timeout=30)
    assert r.status_code == 200
    j = r.json()
    assert "count" in j and "service_count" in j


def test_16_bela_potvrda_label_no_insurance(client):
    """DOC_LABELS should have 'Bela potvrda' without '(osiguranje)'."""
    # Dashboard payload uses DOC_LABELS; excel headers too
    r = client.get(f"{BASE_URL}/api/reports/excel", timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    headers_v = [c.value for c in wb["Vozila"][1]]
    assert "Bela potvrda" in headers_v
    assert not any("Bela potvrda (osiguranje)" == h for h in headers_v)


# ============ CLEANUP ============
def test_zz_cleanup(client):
    if STATE["vehicle_id"]:
        client.delete(f"{BASE_URL}/api/vehicles/{STATE['vehicle_id']}", timeout=30)
    if STATE["driver_id"]:
        client.delete(f"{BASE_URL}/api/drivers/{STATE['driver_id']}", timeout=30)
