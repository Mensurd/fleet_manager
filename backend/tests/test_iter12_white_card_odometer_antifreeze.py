"""Iteration 12 tests: white_card doc, current_odometer endpoint, adjustable antifreeze_every."""
import os
import pytest
import requests
from datetime import date, timedelta
import io as _io
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "http://localhost:8001"
API = f"{BASE_URL}/api"
ADMIN1 = ("duljevic1366@gmail.com", "fleet2026")
ADMIN2 = ("duljevic.pr@gmail.com", "fleet2026")


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=20)
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def hdr():
    return {"Authorization": f"Bearer {_login(*ADMIN1)}"}


# Force sequential run within one xdist worker: use single group
pytestmark = pytest.mark.xdist_group(name="iter12_serial")


def test_both_admin_logins():
    _login(*ADMIN1)
    _login(*ADMIN2)


@pytest.fixture(scope="module")
def throwaway_vehicle(hdr):
    expiring_soon = (date.today() + timedelta(days=15)).isoformat()
    payload = {
        "plate": "TEST-ITR12",
        "make": "TestMake",
        "model": "TestModel",
        "year": 2021,
        "vehicle_type": "Kamion",
        "vin": "TESTVIN12",
        "oil_interval_km": 15000,
        "antifreeze_every": 3,
        "white_card": {"expiry": expiring_soon, "note": "bela"},
    }
    r = requests.post(f"{API}/vehicles", json=payload, headers=hdr, timeout=20)
    assert r.status_code == 200, r.text
    v = r.json()
    yield v
    requests.delete(f"{API}/vehicles/{v['id']}", headers=hdr, timeout=20)


# --------- White card ---------
class TestWhiteCard:
    def test_create_returns_white_card_status(self, throwaway_vehicle):
        v = throwaway_vehicle
        assert "white_card" in v
        assert v["white_card"]["status"] == "expiring"
        assert v["white_card"]["expiry"]

    def test_get_vehicle_has_white_card(self, hdr, throwaway_vehicle):
        r = requests.get(f"{API}/vehicles/{throwaway_vehicle['id']}", headers=hdr)
        assert r.status_code == 200
        assert r.json()["white_card"]["status"] == "expiring"

    def test_dashboard_alerts_include_white_card(self, hdr, throwaway_vehicle):
        r = requests.get(f"{API}/dashboard", headers=hdr)
        assert r.status_code == 200
        alerts = r.json()["alerts"]
        matches = [a for a in alerts if a.get("plate") == "TEST-ITR12" and a.get("doc_key") == "white_card"]
        assert matches, "No white_card alert for throwaway vehicle in dashboard"
        assert matches[0]["doc_type"] == "Bela potvrda (osiguranje)"

    def test_reminders_preview_includes_white_card(self, hdr, throwaway_vehicle):
        # only if in 30/15/7 window - we set 15
        r = requests.get(f"{API}/reminders/preview", headers=hdr)
        assert r.status_code == 200
        items = r.json()["items"]
        assert any(
            i.get("plate") == "TEST-ITR12" and i.get("doc") == "Bela potvrda (osiguranje)"
            for i in items
        ), "White card not in reminders preview"

    def test_excel_export_contains_bela_potvrda_column(self, hdr):
        r = requests.get(f"{API}/reports/excel", headers=hdr)
        assert r.status_code == 200
        wb = load_workbook(_io.BytesIO(r.content), read_only=True)
        ws = wb["Vozila"]
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        assert "Bela potvrda" in header, f"header={header}"

    def test_word_export_contains_bela_potvrda(self, hdr):
        r = requests.get(f"{API}/reports/word", headers=hdr)
        assert r.status_code == 200
        # Docx is zip; check text via python-docx
        from docx import Document
        doc = Document(_io.BytesIO(r.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
        assert "Bela potvrda" in text

    def test_word_per_vehicle_contains_bela_potvrda(self, hdr, throwaway_vehicle):
        r = requests.get(f"{API}/reports/word/{throwaway_vehicle['id']}", headers=hdr)
        assert r.status_code == 200
        from docx import Document
        doc = Document(_io.BytesIO(r.content))
        text = ""
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
        assert "Bela potvrda" in text


# --------- Odometer endpoint ---------
class TestOdometer:
    def test_set_odometer(self, hdr, throwaway_vehicle):
        vid = throwaway_vehicle["id"]
        r = requests.post(f"{API}/vehicles/{vid}/odometer", json={"odometer": 250000}, headers=hdr)
        assert r.status_code == 200, r.text
        # Verify persisted via service-status
        s = requests.get(f"{API}/vehicles/{vid}/service-status", headers=hdr).json()
        assert s["current_odometer"] == 250000
        # remaining = 15000 (base = current since no oil) -> next = 250000+15000, remaining=15000
        assert s["next_oil_odometer"] == 265000
        assert s["remaining_km"] == 15000

    def test_max_of_manual_and_maintenance(self, hdr, throwaway_vehicle):
        vid = throwaway_vehicle["id"]
        # Add a non-oil maintenance with higher odometer
        rm = requests.post(f"{API}/maintenance", json={
            "vehicle_id": vid, "date": "2026-03-01",
            "service_type": "Zamena guma", "description": "tires",
            "cost": 50, "odometer": 300000, "mechanic": "t", "status": "Završeno",
        }, headers=hdr)
        assert rm.status_code == 200
        s = requests.get(f"{API}/vehicles/{vid}/service-status", headers=hdr).json()
        assert s["current_odometer"] == 300000  # max(250000 manual, 300000)


# --------- Antifreeze_every ---------
class TestAntifreezeEvery:
    def test_antifreeze_every_returned(self, hdr, throwaway_vehicle):
        vid = throwaway_vehicle["id"]
        s = requests.get(f"{API}/vehicles/{vid}/service-status", headers=hdr).json()
        assert s["antifreeze_every"] == 3

    def test_antifreeze_due_true_at_multiple_of_three(self, hdr, throwaway_vehicle):
        vid = throwaway_vehicle["id"]
        # Add 2 oil changes so next_change_number = 3 (multiple of 3 -> due True)
        for d, odo in [("2026-04-01", 310000), ("2026-05-01", 325000)]:
            r = requests.post(f"{API}/maintenance", json={
                "vehicle_id": vid, "date": d, "service_type": "Zamena ulja",
                "description": "oil", "cost": 100, "odometer": odo,
                "mechanic": "t", "status": "Završeno",
            }, headers=hdr)
            assert r.status_code == 200
        s = requests.get(f"{API}/vehicles/{vid}/service-status", headers=hdr).json()
        assert s["oil_change_count"] == 2
        assert s["next_change_number"] == 3
        assert s["antifreeze_due_next"] is True

    def test_antifreeze_every_change_to_two_makes_false(self, hdr, throwaway_vehicle):
        vid = throwaway_vehicle["id"]
        # Update antifreeze_every to 2; next=3 -> 3%2=1 -> False
        payload = {
            "plate": "TEST-ITR12", "make": "TestMake", "model": "TestModel",
            "year": 2021, "vehicle_type": "Kamion", "vin": "TESTVIN12",
            "oil_interval_km": 15000, "antifreeze_every": 2,
        }
        r = requests.put(f"{API}/vehicles/{vid}", json=payload, headers=hdr)
        assert r.status_code == 200
        s = requests.get(f"{API}/vehicles/{vid}/service-status", headers=hdr).json()
        assert s["antifreeze_every"] == 2
        assert s["next_change_number"] == 3
        assert s["antifreeze_due_next"] is False
